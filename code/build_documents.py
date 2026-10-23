#!/usr/bin/env python3
from pathlib import Path
import json, math, textwrap
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.section import WD_ORIENT
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[1]; RES=ROOT/'results'; FIG=ROOT/'figures'; DOCS=ROOT/'docs'; EQ=ROOT/'equations'
DOCS.mkdir(exist_ok=True); EQ.mkdir(exist_ok=True)
S=json.loads((RES/'final_summary.json').read_text())
s3=pd.read_csv(RES/'S3_minimal_generator_scores.csv')
s4=pd.read_csv(RES/'S4_dual_EEG_surrogate_tests.csv',keep_default_na=False)
s5=pd.read_csv(RES/'S5_cross_level_reconstructibility.csv')
s6=pd.read_csv(RES/'S6_critical_locus.csv')

def fmt(x, digits=3):
    x=float(x)
    if abs(x)<1e-3 and x!=0: return f"{x:.2e}"
    return f"{x:.{digits}f}"

def equation_png(name, latex, width=7.2):
    p=EQ/f'{name}.png'
    fig=plt.figure(figsize=(width,0.75))
    fig.text(.5,.5,latex,ha='center',va='center',fontsize=18)
    plt.axis('off'); plt.savefig(p,dpi=260,bbox_inches='tight',pad_inches=.08,transparent=True); plt.close(fig)
    return p

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run(); fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '; fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end'); run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)

def base_doc(title_short):
    d=Document(); sec=d.sections[0]; sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.0);sec.right_margin=Cm(2.0)
    styles=d.styles
    st=styles['Normal']; st.font.name='Times New Roman'; st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(5); st.paragraph_format.line_spacing=1.12
    for h,size in [('Title',18),('Heading 1',13),('Heading 2',11.5),('Heading 3',10.5)]:
        styles[h].font.name='Arial'; styles[h].font.size=Pt(size); styles[h].font.bold=True
        if h!='Title': styles[h].font.color.rgb=None
    header=sec.header.paragraphs[0]; header.text=title_short; header.style=styles['Normal']; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; header.runs[0].font.size=Pt(8); header.runs[0].italic=True
    foot=sec.footer.paragraphs[0]; add_page_number(foot)
    return d

def add_title(d,title,subtitle=None):
    p=d.add_paragraph();p.style=d.styles['Title'];p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(title)
    if subtitle:
        q=d.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.LEFT;r=q.add_run(subtitle);r.italic=True;r.font.size=Pt(11)

def add_authors(d):
    p=d.add_paragraph();r=p.add_run('Emirhan Yıldırım');r.bold=True;r.font.size=Pt(11)
    p=d.add_paragraph('Department of Occupational Therapy, Istanbul Medipol University, Istanbul, Türkiye\nFounder / Independent Research Lead, Humanization Tech')
    p.paragraph_format.space_after=Pt(2)
    p=d.add_paragraph('ORCID: 0009-0002-5884-3660 | Correspondence: emirhan.yildirimmy@gmail.com')
    p.runs[0].font.size=Pt(9)

def add_section(d,title,level=1): d.add_heading(title,level=level)
def add_para(d,text,boldlead=None):
    p=d.add_paragraph()
    if boldlead and text.startswith(boldlead):
        r=p.add_run(boldlead);r.bold=True;p.add_run(text[len(boldlead):])
    else:p.add_run(text)
    return p

EQS={
    'eq_orbit': r'$[C]_G=\{Q_A^{T} C Q_B: Q_A\in O(p),Q_B\in O(q)\},\quad \mathcal{Q}_N=X_N/G$',
    'eq_qr': r'$p_r(C)=\mathrm{tr}[(CC^{T})^r],\quad q_r(C)=\frac{p_r(C)}{p_1(C)^r},\quad r\geq2$',
    'eq_erank': r'$r_{\mathrm{eff}}=\exp[-\sum_i\pi_i\log\pi_i],\quad \pi_i=\frac{\sigma_i^2}{\sum_j\sigma_j^2},\quad \mathcal{D}=\{C:\det(CC^T)=0\}$',
    'eq_cross': r'$\Phi:\mathcal{Q}_N\to\mathcal{Q}_S,\quad R_{N\to S}=1-\frac{\mathrm{E}[d(\hat P_S,P_S)^2]}{\mathrm{E}[d(P_S,\bar P_S)^2]}$',
    'sup_orbit': r'$[C]_G=\{Q_A^{T} C Q_B\},\quad k[X_N]^G=\{f:f(g\cdot C)=f(C)\;\forall g\in G\}$',
    'sup_trace': r'$\mathrm{tr}[(C^{\prime}C^{\prime T})^r]=\mathrm{tr}[Q_A^T(CC^T)^rQ_A]=\mathrm{tr}[(CC^T)^r]$',
    'sup_newton': r'$k e_k=\sum_{i=1}^{k}(-1)^{i-1}e_{k-i}p_i,\quad e_0=1$',
    'sup_cca': r'$K=\Sigma_{XX}^{-1/2}\Sigma_{XY}\Sigma_{YY}^{-1/2},\quad \rho_i=\sigma_i(K)$',
}

def add_eq(d,name,text):
    latex=EQS.get(name)
    if latex is None:
        latex='$'+text.replace('_',r'\_')+'$'
    img=equation_png(name,latex)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5)
    p.add_run().add_picture(str(img),width=Inches(5.8))

def add_figure(d,path,caption,width=6.3):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Inches(width))
    c=d.add_paragraph();c.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=c.add_run(caption);r.bold=True;r.font.size=Pt(9)

def add_table(d,headers,rows,widths=None):
    table=d.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.style='Table Grid';hdr=table.rows[0];set_repeat_table_header(hdr);hdr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for j,h in enumerate(headers):
        cell=hdr.cells[j];cell.text=str(h);set_cell_shading(cell,'D9E2F3');cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs:r.bold=True;r.font.size=Pt(8.5)
    for row in rows:
        cells=table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,v in enumerate(row):
            cells[j].text=str(v);cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                for r in p.runs:r.font.size=Pt(8.3)
    if widths:
        for row in table.rows:
            for j,w in enumerate(widths):row.cells[j].width=Cm(w)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

# MAIN MANUSCRIPT
D=base_doc('Relational Identity Under Transformation')
add_title(D,'Relational Identity Under Transformation in Dyadic Neural Systems','A quotient-geometric reformulation of the Intersubjective Contemplative Experiment with synthetic falsification')
add_authors(D)
p=D.add_paragraph();r=p.add_run('Article type: ');r.bold=True;p.add_run('Theory and methods / synthetic falsification study')
p=D.add_paragraph();r=p.add_run('Status: ');r.bold=True;p.add_run('Submission-ready scientific draft; synthetic validation only. No human or animal data were analysed.')

add_section(D,'Abstract')
add_para(D,"Inter-brain synchrony (IBS) is often quantified in coordinate-dependent signal representations, whereas the scientific claim of dyadic identity or shared organization should survive admissible changes of representation. We reformulate the Intersubjective Contemplative Experiment (ICE) as a restricted problem of relational identity under transformation. For a dyadic coupling matrix C, the primary quotient coordinates are normalized spectral power sums q_r = tr[(CC^T)^r]/tr(CC^T)^r, supplemented by rank, effective rank, singular-locus information, a semantic quotient class, and a cross-level reconstructibility map. We then conduct six pre-specified synthetic falsification experiments. Polynomial spectral fingerprints were preserved under O(p) x O(q) basis changes to numerical precision (maximum relative error " + fmt(S['S1']['orthogonal_max_rel_error_max'],2) + ") but changed under anisotropic scaling, while canonical-correlation spectra remained stable under a restricted class of full-rank invertible linear reparameterizations. In six moment-matched relational classes, q2 alone produced " + fmt(100*S['S3']['scores']['q2']['ood'],1) + "% out-of-distribution (OOD) reconstruction, q2+q3 reached " + fmt(100*S['S3']['scores']['q2_q3']['ood'],1) + "%, whereas raw coordinates fell to " + fmt(100*S['S3']['scores']['raw_coordinates']['ood'],1) + "% (chance 16.7%). In EEG-like simulations, partner-specific alpha coupling survived a dyad-shuffled surrogate test (mean PLV delta=" + fmt(S['S4']['partner_alpha_plv']['delta'],4) + ", 95% bootstrap CI [" + fmt(S['S4']['partner_alpha_plv']['ci_low'],4) + ", " + fmt(S['S4']['partner_alpha_plv']['ci_high'],4) + "], p=" + fmt(S['S4']['partner_alpha_plv']['signflip_p'],5) + "), but the q2 summary of the PLV matrix did not. Cross-level invariant features reconstructed a deliberately linked synthetic semantic target (mean OOD R²=" + fmt(S['S5']['linked']['invariant_ood_r2'],3) + ") yet also retained non-zero predictability under a common-cause world (R²=" + fmt(S['S5']['common_cause']['invariant_ood_r2'],3) + "), falsifying any causal interpretation of reconstructibility alone. Finally, effective rank declined smoothly before exact rank loss at the algebraic singular locus, separating an early capacity descriptor from a discrete identity-boundary claim. The framework therefore supports transformation-specific relational fingerprints, not claims of shared consciousness or empirical intentionality-driven IBS. It yields a falsifiable, surrogate-controlled path for prospective dual-EEG testing.")
p=D.add_paragraph();r=p.add_run('Keywords: ');r.bold=True;p.add_run('algebraic invariants; quotient geometry; hyperscanning; inter-brain synchrony; representational similarity; reconstructibility; effective rank; synthetic validation')

add_section(D,'1. Introduction')
add_para(D,"Hyperscanning extends neuroscience from single participants to simultaneously recorded dyads or groups. Yet a central inferential problem remains: an inter-brain statistic may change because the underlying relation changed, because the measurement coordinate system changed, or because both participants were driven by the same task, stimulus, preprocessing artifact, or spectral structure. Phase-based hyperscanning measures are especially sensitive to analysis choices, and pseudo-dyad or dyad-shuffled controls are therefore essential for distinguishing partner-specific coupling from parallel task engagement [2-6].")
add_para(D,"The original ICE manuscript proposed an extreme experimental paradigm in which shared sensory and motor cues are minimized while a shared intentional target is retained. Its strongest empirical proposal was prospective dual-EEG hyperscanning, with IBS related to independently coded semantic convergence. The present work removes claims that are not required for that test. Empathy, meditation, frequency bands, and symbolic tokens are treated as moderators, operators, measurement channels, or observables rather than invariants. The theory is reduced to a narrower question: which relational structures remain identifiable under admissible transformations, and which of those structures are sufficient to reconstruct a dyadic identity class?")
add_para(D,"This shift follows a general invariant-theoretic principle: coordinates are not the object of interest when a transformation group acts on the representation. The scientific object is the orbit or quotient class, together with a minimal set of functions that separates the classes relevant to the experiment. Because maximal invariance can itself erase meaningful differences, the transformation group must be declared rather than assumed. This is particularly important for invertible-linear invariance in finite-sample representational comparison [7].")

add_section(D,'2. Quotient-geometric formulation')
add_section(D,'2.1 Dyadic state, group action, and quotient',2)
add_para(D,"Let C be a p x q dyadic relation matrix derived from two simultaneously observed systems. C may represent a cross-covariance, cross-spectral, PLV, or another explicitly defined coupling operator. Under independent orthogonal changes of basis, C transforms as C' = Q_A^T C Q_B, where Q_A in O(p) and Q_B in O(q). The individual matrix entries are representation-dependent; the orbit [C]_G is the object whose identity is preserved.")
add_eq(D,'eq_orbit','[C]G = {Q_A^T C Q_B : Q_A ∈ O(p), Q_B ∈ O(q)},    Q_N = X_N // G')
add_section(D,'2.2 Spectral polynomial fingerprint',2)
add_para(D,"Define M_C = CC^T. Independent orthogonal basis changes induce a similarity transformation of M_C, so its eigenvalue multiset and all symmetric polynomials in those eigenvalues are invariant. We use power sums p_r and their scale-normalized forms q_r.")
add_eq(D,'eq_qr','p_r(C) = tr[(CC^T)^r],    q_r(C) = p_r(C) / p_1(C)^r,    r ≥ 2')
add_para(D,"The characteristic polynomial coefficients are equivalent spectral coordinates. In finite dimension d, p_1,...,p_d determine those coefficients through Newton identities, so a complete spectral fingerprint exists in principle. The practical question is smaller: which subset is sufficient for the class family under study?")
add_section(D,'2.3 Capacity and singular strata',2)
add_para(D,"Effective rank is used as a continuous spectral-capacity descriptor, not as a discrete algebraic rank. The exact rank-deficient locus is defined separately. This prevents an arbitrary effective-rank threshold from being mislabelled as an identity boundary.")
add_eq(D,'eq_erank','r_eff(C) = exp[-Σ_i π_i log π_i],    π_i = σ_i² / Σ_j σ_j²;    D = {C : det(CC^T) = 0}')
add_section(D,'2.4 Semantic quotient and cross-level reconstructibility',2)
add_para(D,"Raw words or embedding coordinates are not semantic invariants. Let H act on a semantic representation space S; the semantic object is an equivalence class [y]_H in Q_S = S//H. A cross-level hypothesis is then a map Phi from neural quotient coordinates to semantic quotient coordinates. Its empirical analogue is out-of-distribution reconstructibility. Crucially, reconstructibility is an association property and does not establish causal direction.")
add_eq(D,'eq_cross','Φ : Q_N → Q_S,    R_N→S = 1 - E[d(P̂_S, P_S)²] / E[d(P_S, P̄_S)²]')

add_section(D,'3. Synthetic falsification study')
add_para(D,"All experiments were generated de novo with a fixed master seed (20260812) and are fully reproducible from the accompanying Python script. No empirical neural data were used. The study was designed around negative controls and failure modes rather than confirmation alone.")
add_table(D,['Experiment','Primary question','Positive control','Negative/boundary control'],[
('S1','Are q_r exact under the declared O(p) x O(q) action?','Random orthogonal reparameterization','Anisotropic non-orthogonal scaling'),
('S2','Can a broader linear-invariance candidate be defined?','CCA spectrum under invertible transforms','q2 under the same transforms'),
('S3','Which fingerprint components reconstruct relational class OOD?','Moment-matched spectral classes','Raw-coordinate OOD baseline; ablation'),
('S4','Can partner specificity be separated from shared input?','Dyad-specific alpha phase trajectories','Partner shuffle; global theta common input'),
('S5','Does neural fingerprint reconstruct a semantic target?','Explicit linked world','Independent null; common-cause world; shuffled labels'),
('S6','How does the system approach rank loss?','Controlled spectral collapse','Exact singular endpoint vs continuous effective rank')],widths=[1.3,4.6,4.2,4.4])

add_section(D,'4. Results')
add_section(D,'4.1 S1: transformation-specific invariance',2)
add_para(D,"Across 1,000 random 8 x 8 coupling matrices, the maximum relative discrepancy in the normalized power-sum fingerprint after independent orthogonal basis changes was " + fmt(S['S1']['orthogonal_max_rel_error_max'],2) + ", consistent with floating-point error. In contrast, anisotropic scaling produced a median log-fingerprint distance of " + fmt(S['S1']['anisotropic_distance_median'],3) + " (95th percentile " + fmt(S['S1']['anisotropic_distance_p95'],3) + "). Thus the present q_r family is explicitly O(p) x O(q)-invariant, not GL(p) x GL(q)-invariant.")
add_figure(D,FIG/'Figure_1_Transformation_Invariance.png','Figure 1. Exact preservation under the declared orthogonal action and deliberate failure outside that group.')

add_section(D,'4.2 S2: restricted invertible-linear extension',2)
add_para(D,"In 300 full-rank latent-Gaussian simulations (n=600, p=q=6), the canonical-correlation spectrum was preserved after independently generated invertible linear transformations with maximum absolute discrepancy " + fmt(S['S2']['cca_max_abs_error_max'],2) + ". Under the same transforms, q2 changed by a median relative amount of " + fmt(S['S2']['q2_relative_change_median'],3) + ". This establishes a restricted GL candidate but not a universal representation metric; high-dimensional invertible invariance can become non-informative when representation dimension approaches or exceeds sample size [7].")

add_section(D,'4.3 S3: minimal-generator and reconstructibility test',2)
rows=[]
for _,r in s3.iterrows(): rows.append((r.feature_set,fmt(100*r.train_accuracy,1)+'%',fmt(100*r.ood_accuracy,1)+'%'))
add_table(D,['Fingerprint','Training accuracy','OOD accuracy'],rows,widths=[5.5,4,4])
add_para(D,"The six classes were selected to have nearly matched q2 but differing higher spectral moments. q2 alone yielded " + fmt(100*S['S3']['scores']['q2']['ood'],1) + "% OOD accuracy. Adding q3 increased OOD reconstruction to " + fmt(100*S['S3']['scores']['q2_q3']['ood'],1) + "%. The raw-coordinate classifier, despite perfect training accuracy, fell to " + fmt(100*S['S3']['scores']['raw_coordinates']['ood'],1) + "%, close to the 16.7% six-class chance level after unseen rotations. In this deliberately constructed family, q2+q3 is sufficient; this is a family-specific minimal generating fingerprint, not a universal theorem. Using p1,...,p8 and Newton identities reconstructed the eight eigenvalues with median maximum numerical error " + fmt(S['S3']['newton_median_max_error'],2) + ".")
add_figure(D,FIG/'Figure_2_Minimal_Generators.png','Figure 2. OOD reconstruction after invariant-generator ablation. Raw coordinates collapse after unseen basis changes.')

add_section(D,'4.4 S4: EEG-like partner-shuffled surrogate test',2)
pa=S['S4']['partner_alpha_plv']; pq=S['S4']['partner_alpha_q2']; ct=S['S4']['common_theta_plv']
add_para(D,"The EEG-like simulation used 48 dyads per world, six channels per participant, 8 s at 128 Hz, 1/f-like AR noise, and band-limited theta, alpha, and beta components. In the partner-specific world, paired participants shared a time-varying alpha phase trajectory unique to that dyad. True-pair mean alpha PLV exceeded the partner-shuffled surrogate by " + fmt(pa['delta'],4) + " (95% bootstrap CI [" + fmt(pa['ci_low'],4) + ", " + fmt(pa['ci_high'],4) + "], sign-flip p=" + fmt(pa['signflip_p'],5) + "). In contrast, the q2 summary of the same PLV matrices did not distinguish true from shuffled partners (delta=" + fmt(pq['delta'],2) + ", p=" + fmt(pq['signflip_p'],3) + "). A global common-input theta trajectory also failed partner specificity (mean PLV delta=" + fmt(ct['delta'],4) + ", p=" + fmt(ct['signflip_p'],3) + "). Thus q2 is not a universal partner-specificity statistic, and a high PLV must be tested against dyad-shuffled surrogates.")
add_figure(D,FIG/'Figure_3_EEG_Surrogate_Controls.png','Figure 3. Partner-shuffled controls separate a dyad-specific alpha trajectory from a common-input theta trajectory.')

add_section(D,'4.5 S5: cross-level reconstructibility and causal failure',2)
linked=S['S5']['linked']; indep=S['S5']['independent']; cc=S['S5']['common_cause']
add_para(D,"When the synthetic semantic target was explicitly generated from invariant fingerprint coordinates, the invariant model generalized across unseen orthogonal reparameterizations (mean OOD R²=" + fmt(linked['invariant_ood_r2'],3) + "), whereas the coordinate-sensitive raw baseline failed strongly under the shift. The independent-null world produced mean invariant OOD R²=" + fmt(indep['invariant_ood_r2'],3) + ". However, a latent common cause that jointly changed relational spectra and semantic output produced positive reconstructibility (R²=" + fmt(cc['invariant_ood_r2'],3) + ") despite no imposed neural-to-semantic causal map. Cross-level reconstructibility therefore survives as an association criterion only.")
add_figure(D,FIG/'Figure_4_Cross_Level_Confound.png','Figure 4. Cross-level reconstructibility succeeds in a linked world but remains vulnerable to common-cause structure.')

add_section(D,'4.6 S6: critical locus and the distinction between effective and exact rank',2)
add_para(D,"Along a controlled spectral-collapse path, mean effective rank declined from " + fmt(S['S6']['0.0']['effective_rank'],3) + " at t=0 to " + fmt(S['S6']['0.95']['effective_rank'],3) + " at t=0.95 while exact matrix rank remained 8. Only at the constructed endpoint t=1 did exact rank become 1 and the determinant locus become singular. Thus effective rank is a graded capacity marker; it is not by itself an algebraic phase boundary.")
add_figure(D,FIG/'Figure_5_Critical_Locus.png','Figure 5. Continuous loss of effective capacity precedes the exact rank-deficient algebraic locus.')

add_section(D,'5. Discussion')
add_para(D,"The synthetic study narrows the theory in six ways. First, the robust object is not an EEG amplitude, frequency band, word, or raw PLV matrix; it is an explicitly transformation-indexed relational class. Second, normalized trace moments provide exact polynomial invariants under orthogonal basis changes, but the admissible group cannot be silently enlarged. Third, full-rank CCA spectra offer a restricted invertible-linear extension, subject to finite-sample and dimensionality limitations. Fourth, no single low-order moment is guaranteed to identify a relational class: generator sufficiency depends on the family of alternatives that must be separated. Fifth, partner specificity is a separate statistical question from basis invariance; the EEG-like experiment showed that mean alpha PLV could carry partner-specific information while q2 of the PLV matrix did not. Sixth, cross-level reconstructibility is not causality, and effective-rank decline is not equivalent to crossing a singular stratum.")
add_para(D,"These distinctions are important for ICE. The scientifically defensible core is not a claim that intention remotely synchronizes brains, nor that semantic similarity demonstrates a shared conscious state. The defensible core is a prospective falsification problem: after removing shared sensory/motor channels as far as feasible, does a pre-registered relational fingerprint show true-dyad structure beyond pseudo-dyad controls, and does any independently scored semantic correspondence add out-of-sample information without being explainable by common timing, task structure, signal quality, or other shared causes?")
add_para(D,"Current hyperscanning methodology supports this conservative stance. Dual-EEG pipelines explicitly include controls for spurious correlations [2]; preprocessing choices can materially change PLV estimates [3]; methodological degrees of freedom can bias IBS estimates [4]; and recent work distinguishes task-enhanced synchrony from partner-specific synchrony using dyad-shuffled surrogate controls and spectral parameterization [5]. These considerations are methodological constraints on any future ICE test, not evidence for ICE itself.")

add_section(D,'6. Prospective empirical protocol')
add_para(D,"A future empirical study should be preregistered before inspection of outcomes. The primary comparison should be true dyads versus partner-shuffled pseudo-dyads within matched condition, not simply ICE versus rest. A factorial design should separate target matching (same vs different internally instructed target), interaction status (paired protocol vs independent control), and physical/sensory coupling (isolated vs an explicitly shared-input positive control). The analysis should predefine preprocessing, epoch length, artifact rejection, spectral parameterization, and all surrogate constructions.")
add_table(D,['Domain','Primary preregistered requirement'],[
('Neural','Dual EEG; synchronized clocks; predeclared preprocessing; PLV plus at least one complementary phase/correlation estimator; spectral/aperiodic characterization.'),
('Partner specificity','True dyad compared with many partner-shuffled pseudo-dyads matched for condition and timing.'),
('Invariant fingerprint','Declare C, admissible group G, q_r set, rank/effective-rank descriptors, and any CCA/whitening extension before testing.'),
('Semantics','Blind, independent scoring or locked embedding pipeline; no post-hoc symbolic ontology; permutation nulls.'),
('Cross-level','Nested/held-out prediction; common-cause sensitivity models; no causal language from R²/correlation alone.'),
('Falsification','Predefined failure gates: no partner-specific effect; effect unstable to preprocessing; semantic effect absent OOD; invariant breaks under declared group.')],widths=[3.2,11.5])

add_section(D,'7. Limitations')
add_para(D,"All numerical results in this article are synthetic. They validate algebraic and statistical behavior of candidate constructions under specified generators; they do not establish a biological invariant, inter-brain mechanism, intentionality-driven coupling, empathy-mediated causation, semantic transmission, shared consciousness, or clinical relevance. The EEG simulator is intentionally simplified and does not reproduce the full spatial mixing, referencing, non-stationarity, muscle/ocular artifacts, volume conduction, or aperiodic structure of real dual-EEG recordings. The semantic quotient remains underspecified until an empirical representation space and admissible semantic transformation group are fixed. The CCA extension assumes well-conditioned covariance estimates and sufficient sample size. Finally, a complete invariant ring can be statistically useless, while a statistically predictive summary can fail to be invariant; both algebraic and empirical criteria are therefore required.")

add_section(D,'8. Conclusion')
add_para(D,"ICE can be reduced from a broad neuro-phenomenological theory to a falsifiable relational-identity problem. The present synthetic study supports a transformation-specific quotient geometry in which spectral power sums provide exact O(p) x O(q) invariants, higher moments may be necessary for class reconstruction, restricted CCA spectra extend invariance under full-rank linear reparameterizations, partner-shuffled surrogates are indispensable for hyperscanning inference, and cross-level reconstructibility is explicitly non-causal. The next scientific step is not further theoretical expansion but a preregistered empirical dual-EEG study designed to falsify these restricted claims.")

add_section(D,'Declarations')
add_para(D,'Data and code availability. All data in the present study are synthetic and are regenerated by the accompanying script. CSV outputs, figures, the master seed, environment requirements, and SHA-256 manifest are included in the submission package.')
add_para(D,'Ethics. No human participants, animals, identifiable private data, or clinical records were used in this synthetic study; institutional ethics review was therefore not applicable to the analyses reported here. Any prospective ICE hyperscanning study will require independent ethics review and informed consent before data collection.')
add_para(D,'Funding. No external funding is declared for this synthetic study.')
add_para(D,'Competing interests. The author is Founder / Independent Research Lead at Humanization Tech. No commercial or clinical claim is made from the synthetic results.')
add_para(D,'Author contributions (CRediT). Emirhan Yıldırım: Conceptualization, Methodology, Formal analysis, Investigation, Software, Visualization, Writing - original draft, Writing - review and editing.')

add_section(D,'References')
refs=[
"1. Yıldırım E. The Intersubjective Contemplative Experiment (ICE): A Neuro-Phenomenological Framework for Investigating Shared Experience via Pure Intentionality. Conceptual manuscript, 2026.",
"2. Kayhan E, Matthes D, Haresign IM, et al. DEEP: A dual EEG pipeline for developmental hyperscanning studies. Developmental Cognitive Neuroscience. 2022;54:101104. doi:10.1016/j.dcn.2022.101104.",
"3. Pili MP, Provenzi L, Billeci L, et al. Exploring the impact of manual and automatic EEG pre-processing methods on interpersonal neural synchrony measures in parent-infant hyperscanning studies. Journal of Neuroscience Methods. 2025;417:110400. doi:10.1016/j.jneumeth.2025.110400.",
"4. Zimmermann M, Schultz-Nielsen K, Dumas G, Konvalinka I. Arbitrary methodological decisions skew inter-brain synchronization estimates in hyperscanning-EEG studies. Imaging Neuroscience. 2024;2:1-19. doi:10.1162/imag_a_00350.",
"5. Valenzuela L, Cavallaro R, Mangeng A, et al. Aperiodic activity shapes inter-brain synchrony during cooperative puzzle solving: dissociable task-enhanced and partner-specific spectral signatures. NeuroImage. 2026:122116. doi:10.1016/j.neuroimage.2026.122116.",
"6. Burgess AP. On the interpretation of synchronization in EEG hyperscanning studies: a cautionary note. Frontiers in Human Neuroscience. 2013;7:881. doi:10.3389/fnhum.2013.00881.",
"7. Kornblith S, Norouzi M, Lee H, Hinton G. Similarity of Neural Network Representations Revisited. Proceedings of Machine Learning Research. 2019;97:3519-3529.",
"8. Hakim U, De Felice S, Pinti P, et al. Quantification of inter-brain coupling: A review of current methods used in haemodynamic and electrophysiological hyperscanning studies. NeuroImage. 2023;280:120354. doi:10.1016/j.neuroimage.2023.120354.",
"9. Dowdall JR, Vinck M. Coherence fails to reliably capture inter-areal interactions in bidirectional neural systems with transmission delays. NeuroImage. 2023;271:119998. doi:10.1016/j.neuroimage.2023.119998.",
"10. Weyl H. The Classical Groups: Their Invariants and Representations. Princeton University Press.",
"11. Mumford D, Fogarty J, Kirwan F. Geometric Invariant Theory. 3rd ed. Springer; 1994."
]
for ref in refs:
    p=D.add_paragraph(ref); p.paragraph_format.left_indent=Cm(.4); p.paragraph_format.first_line_indent=Cm(-.4); p.paragraph_format.space_after=Pt(3)

main_path=DOCS/'ICE_Quotient_Geometric_Invariants_Main_Manuscript.docx';D.save(main_path)

# SUPPLEMENT
A=base_doc('Supplementary Appendix - ICE Quotient Geometry')
add_title(A,'Supplementary Appendix','Formal definitions, proofs, synthetic generators, statistical controls, and reproducibility ledger')
add_authors(A)
add_para(A,'This supplement accompanies the manuscript “Relational Identity Under Transformation in Dyadic Neural Systems.” All results are synthetic. The supplement is intentionally explicit about negative controls and claim boundaries.')
add_section(A,'Appendix A. Formal definitions and algebraic claims')
add_section(A,'A.1 Transformation-indexed identity',2)
add_para(A,'Let X_N be a dyadic relation space and G = O(p) x O(q) act on C by (Q_A,Q_B)·C = Q_A^T C Q_B. The relational identity object is the orbit [C]_G. A scalar f is a G-invariant if f(g·C)=f(C) for every g in G. This definition always includes the group; “invariant” without a declared transformation class is incomplete.')
add_eq(A,'sup_orbit','[C]G = {Q_A^T C Q_B},    k[X_N]^G = {f : f(g·C)=f(C) for all g∈G}')
add_section(A,'A.2 Proof of the spectral power-sum invariants',2)
add_para(A,'For C\' = Q_A^T C Q_B with Q_A and Q_B orthogonal, C\'C\'^T = Q_A^T CC^T Q_A. The matrices CC^T and C\'C\'^T are similar and therefore have identical characteristic polynomials and eigenvalue multisets. Cyclic invariance of trace gives p_r(C\')=p_r(C) for every positive integer r. Because p_1 is also invariant, q_r=p_r/p_1^r is scale-normalized and O(p) x O(q)-invariant wherever p_1>0.')
add_eq(A,'sup_trace','tr[(C′C′^T)^r] = tr[Q_A^T (CC^T)^r Q_A] = tr[(CC^T)^r]')
add_section(A,'A.3 Newton completeness of the finite spectral fingerprint',2)
add_para(A,'Let lambda_1,...,lambda_d denote eigenvalues of CC^T. Power sums p_k=sum_i lambda_i^k and elementary symmetric polynomials e_k are related by Newton identities. Thus p_1,...,p_d determine the characteristic polynomial coefficients and hence the eigenvalue multiset, up to numerical root finding. The numerical experiment is a stability check, not a proof of Newton identities.')
add_eq(A,'sup_newton','k e_k = Σ_{i=1}^k (-1)^(i-1) e_{k-i} p_i,    e_0 = 1')
add_section(A,'A.4 Restricted GL extension via canonical correlations',2)
add_para(A,'For centered observations X and Y with positive-definite within-system covariances, canonical correlations are singular values of the whitened cross-covariance K = Sigma_XX^{-1/2} Sigma_XY Sigma_YY^{-1/2}. They are invariant to nonsingular linear reparameterizations in the classical full-rank setting. This should not be interpreted as a universal preference for GL invariance; with high-dimensional representations, excessive invariance can destroy meaningful similarity information [Kornblith et al., 2019].')
add_eq(A,'sup_cca','K = Σ_XX^(-1/2) Σ_XY Σ_YY^(-1/2),    ρ_i = σ_i(K)')
add_section(A,'A.5 Singular strata and capacity',2)
add_para(A,'The algebraic rank-deficient locus is D={C:det(CC^T)=0}. Effective rank is continuous for positive singular values and is used only as a graded capacity descriptor. In the synthetic collapse, exact rank remains unchanged until a singular value becomes exactly zero.')
add_section(A,'A.6 Semantic quotient',2)
add_para(A,'No semantic embedding is declared invariant a priori. If a representation y is acted upon by a declared group H, the object is [y]_H. Pairwise cosine similarity is only invariant under common isometries and therefore cannot serve as a general semantic quotient coordinate under arbitrary embedding changes. The empirical study must freeze H or use a relational semantic representation with explicit invariance properties.')

add_section(A,'Appendix B. Synthetic experiment specifications')
add_section(A,'B.1 S1 - orthogonal invariance and out-of-group scaling',2)
add_para(A,'Generator: 1,000 independent 8 x 8 Gaussian matrices normalized to unit Frobenius norm. Positive transformation: independent Haar-like QR orthogonal matrices. Negative transformation: independent diagonal log-normal anisotropic scalings. Endpoint metrics: maximum relative q_1...q_4 discrepancy and Euclidean distance in log fingerprint coordinates.')
add_table(A,['Metric','Result'],[
('Maximum O x O relative error',fmt(S['S1']['orthogonal_max_rel_error_max'],2)),('Median O x O relative error',fmt(S['S1']['orthogonal_max_rel_error_median'],2)),('Median out-of-group log distance',fmt(S['S1']['anisotropic_distance_median'],4)),('95th percentile out-of-group distance',fmt(S['S1']['anisotropic_distance_p95'],4))])
add_section(A,'B.2 S2 - restricted GL extension',2)
add_para(A,'For each of 300 replicates, n=600 samples were generated from three latent factors into p=q=6 observed dimensions with independent Gaussian noise. X and Y were independently transformed by random well-conditioned invertible matrices. The CCA spectrum was compared before and after transformation; q2 of normalized raw cross-covariance was the boundary comparator.')
add_table(A,['Metric','Result'],[
('Maximum CCA-spectrum absolute error',fmt(S['S2']['cca_max_abs_error_max'],2)),('Median CCA-spectrum absolute error',fmt(S['S2']['cca_max_abs_error_median'],2)),('Median q2 relative change',fmt(S['S2']['q2_relative_change_median'],4))])
add_figure(A,FIG/'Supplementary_Figure_S1_GL_Extension.png','Supplementary Figure S1. CCA-spectrum invariance versus q2 change under invertible linear reparameterization.')

add_section(A,'B.3 S3 - moment-matched relational classes',2)
add_para(A,'Candidate 8-component spectral distributions were drawn from a Dirichlet family and filtered to a narrow q2 interval. Six templates were selected by farthest-point sampling in (q3,q4) space. Each trial received multiplicative spectral jitter and was normalized. Training matrices were near-canonical; OOD test matrices received independent random orthogonal left/right rotations. Logistic classifiers were fitted to progressively larger invariant subsets and to flattened raw coordinates.')
add_table(A,['Fingerprint','Train','OOD'],[(r.feature_set,fmt(r.train_accuracy,4),fmt(r.ood_accuracy,4)) for _,r in s3.iterrows()])
add_para(A,'Interpretation: q2 is deliberately non-separating. q2+q3 is sufficient for this six-class family; q4 and effective rank are redundant here. The result does not imply that q3 is universally sufficient. The finite-dimensional complete spectrum can be recovered from enough power sums; in the numerical Newton check the median maximum eigenvalue error was '+fmt(S['S3']['newton_median_max_error'],2)+' and the worst observed error was '+fmt(S['S3']['newton_worst_max_error'],2)+'.')

add_section(A,'B.4 S4 - EEG-like dual-signal surrogate world',2)
add_para(A,'Each synthetic world contained 48 dyads, six channels per participant, 8 s per dyad at 128 Hz. Signals contained AR(1)-like low-frequency background plus theta/alpha/beta oscillatory components with channel mixing. Partner-specific alpha was generated by a unique time-varying phase trajectory shared only within the true dyad. Common-input theta used one trajectory shared across all dyads. Partners were then circularly shifted across dyads to create pseudo-dyads. PLV matrices were estimated after fourth-order Butterworth band-pass filtering and Hilbert phase extraction. Paired differences used bootstrap confidence intervals and sign-flip tests.')
selected=s4[(s4.world.isin(['partner_specific','common_input','independent_null'])) & (s4.band.isin(['theta','alpha']))]
world_label={'partner_specific':'partner specific','common_input':'common input','independent_null':'independent null'}
add_table(A,['World','Band','Metric','Real','Shuffled','Delta','95% CI','p'],[(world_label.get(r.world,r.world),r.band,r.metric,fmt(r.real_mean,4),fmt(r.shuffled_mean,4),fmt(r.delta,4),f"[{fmt(r.ci_low,4)}, {fmt(r.ci_high,4)}]",fmt(r.signflip_p,5)) for _,r in selected.iterrows()],widths=[2.7,1.5,2.6,1.7,1.7,1.7,3.1,1.5])
add_para(A,'Predeclared falsification reading: partner-specific alpha mean PLV passes the pseudo-dyad control in this generator; q2 of its PLV matrix does not. Common-input theta does not pass partner specificity. Therefore “algebraically invariant summary” and “statistically partner-specific observable” are distinct properties.')

add_section(A,'B.5 S5 - cross-level reconstructibility',2)
add_para(A,'Three worlds were generated in 15 Monte Carlo replicates with 600 training and 300 OOD test trials per replicate. Invariant predictors were q2, q3, q4 and effective rank. OOD matrices received unseen orthogonal rotations. Linked: the semantic target was explicitly generated from invariant coordinates. Independent: targets were independent Gaussian noise. Common cause: a latent z influenced both the relational spectrum and semantic outcome, without an imposed structural neural-to-semantic map. A polynomial ridge model was compared with a raw-coordinate ridge model and a shuffled-label control.')
g=s5.groupby('world')[['invariant_ood_r2','raw_ood_r2','trial_shuffled_r2']].agg(['mean','std'])
rows=[]
for world in ['linked','independent','common_cause']:
    sub=s5[s5.world==world]
    rows.append((world.replace('_',' '),fmt(sub.invariant_ood_r2.mean(),4)+' ± '+fmt(sub.invariant_ood_r2.std(),4),fmt(sub.raw_ood_r2.mean(),4)+' ± '+fmt(sub.raw_ood_r2.std(),4),fmt(sub.trial_shuffled_r2.mean(),4)+' ± '+fmt(sub.trial_shuffled_r2.std(),4)))
add_table(A,['World','Invariant OOD R²','Raw OOD R²','Shuffled-label OOD R²'],rows,widths=[3.1,4,4,4])
add_para(A,'The common-cause result is a deliberate causal falsification. Positive cross-level R² is compatible with a common driver and is therefore insufficient for causal language.')

add_section(A,'B.6 S6 - controlled approach to a singular locus',2)
add_para(A,'The singular-value path was s(t)=(1,(1-t)0.9,...,(1-t)0.3). Random left/right orthogonal rotations were applied. One hundred replicates were evaluated at each t. Effective rank, smallest singular value, exact numerical rank and log determinant were recorded.')
for tt in [0,.5,.9,.95,.99,1.0]:
    sub=s6[np.isclose(s6.collapse_t,tt)]
    if len(sub):
        pass
add_table(A,['t','Mean effective rank','Mean min singular value','Exact rank'],[(fmt(tt,2),fmt(sub.effective_rank.mean(),4),fmt(sub.minimum_singular_value.mean(),5),str(int(round(sub.exact_rank.mean())))) for tt in sorted(s6.collapse_t.unique()) for sub in [s6[np.isclose(s6.collapse_t,tt)]]],widths=[2,4,4,3])

add_section(A,'Appendix C. Statistical specification and robustness rules')
add_para(A,'C.1 All random generation uses NumPy Generator with master seed 20260812. The pipeline does not select seeds based on performance.')
add_para(A,'C.2 Partner-shuffled EEG tests are paired at the A-participant level. The null is generated by mapping each A participant to the next B participant within the same synthetic condition. Confidence intervals resample paired differences. Two-sided sign-flip p-values include a +1 correction in numerator and denominator.')
add_para(A,'C.3 OOD means a transformation distribution absent from training. It does not mean biological or population OOD generalization.')
add_para(A,'C.4 No p-value is used to claim a biological mechanism. S4 p-values validate recovery of signals deliberately inserted into the simulator.')
add_para(A,'C.5 No threshold on effective rank is declared an algebraic boundary. Exact rank loss and det(CC^T)=0 define the modeled singular stratum.')

A.add_page_break()
add_section(A,'Appendix D. Falsification and claim-audit ledger')
add_table(A,['Candidate claim','Synthetic verdict','Permitted manuscript claim','Prohibited inference'],[
('q_r is invariant to arbitrary coordinate transforms','FALSIFIED','q_r is exact under O(p) x O(q)','Do not call q_r GL-invariant'),
('A broader linear invariant exists','SUPPORTED in restricted full-rank setting','CCA spectrum is stable under tested invertible transforms','Do not claim universal GL similarity'),
('q2 alone identifies relational identity','FALSIFIED','Higher moments may be necessary','Do not use q2 as universal fingerprint'),
('Algebraic q2 of PLV matrix detects partner specificity','FALSIFIED in S4','Mean alpha PLV recovered inserted partner-specific trajectory','Do not conflate invariance with partner specificity'),
('Cross-level reconstructibility implies causation','FALSIFIED','Reconstructibility is an association criterion','Do not infer neural -> semantic causality from R²'),
('Effective-rank threshold is identity boundary','FALSIFIED','Effective rank is graded; exact singular locus is discrete','Do not choose post-hoc threshold and call it algebraic'),
('Synthetic study validates ICE biologically','NOT TESTED','Framework is ready for prospective empirical falsification','No shared-consciousness or remote-IBS claim')],widths=[4,2.6,5,5])

add_section(A,'Appendix E. Reproducibility map')
add_para(A,'Primary executable: code/ice_invariant_validation.py. It regenerates all CSV outputs, final_summary.json and six figures. Document builder: code/build_documents.py. Requirements are pinned at the package level. The top-level run_all.sh executes both stages. The SHA-256 manifest permits verification of delivered files.')
add_table(A,['Output','Purpose'],[
('results/S1_transformation_invariance.csv','Group-action and out-of-group tests'),('results/S2_GL_extension.csv','Restricted CCA/GL test'),('results/S3_minimal_generator_scores.csv','Generator-ablation/OOD classification'),('results/S3_moment_matched_templates.csv','Synthetic spectral class templates'),('results/S3_newton_spectrum_reconstruction.csv','Power-sum completeness numerical check'),('results/S4_dual_EEG_pairwise_metrics.csv','Dyad-level real and shuffled EEG-like metrics'),('results/S4_dual_EEG_surrogate_tests.csv','Bootstrap/sign-flip surrogate summaries'),('results/S5_cross_level_reconstructibility.csv','Linked/null/common-cause Monte Carlo R²'),('results/S6_critical_locus.csv','Rank-collapse path'),('results/final_summary.json','Machine-readable headline results')],widths=[7.2,8])

add_para(A,'References are numbered and provided in full in the main manuscript.')
supp_path=DOCS/'ICE_Quotient_Geometric_Invariants_Supplementary_Appendix.docx';A.save(supp_path)

# Neutral cover letter draft
C=base_doc('Cover Letter')
add_title(C,'Cover Letter - Journal Submission Draft')
p=C.add_paragraph('Dear Editor,')
add_para(C,'Please consider our manuscript, “Relational Identity Under Transformation in Dyadic Neural Systems: A quotient-geometric reformulation of the Intersubjective Contemplative Experiment with synthetic falsification,” as a theory-and-methods contribution.')
add_para(C,'The manuscript addresses a methodological problem in hyperscanning and representational neuroscience: how to distinguish genuine relational structure from changes caused by coordinate choice, shared inputs, and analytic representation. We derive transformation-indexed spectral invariants, test restricted extensions under invertible linear reparameterization, and evaluate the framework through six reproducible synthetic falsification experiments including pseudo-dyad controls, OOD reconstruction, minimal-generator ablation, common-cause confounding, and a controlled singular-locus experiment.')
add_para(C,'The manuscript is intentionally claim-bounded. It does not report human or animal data and does not claim that shared intentionality causes inter-brain synchrony. Instead, it converts a broad conceptual proposal into a falsifiable mathematical and statistical framework and identifies several negative results that constrain future empirical use.')
add_para(C,'All synthetic data, code, figures, and a detailed supplementary appendix are included in the reproducibility package. The work is original to the best of the author’s knowledge and is not presented as clinical validation.')
C.add_paragraph('Sincerely,\nEmirhan Yıldırım\nIstanbul Medipol University / Humanization Tech')
C.save(DOCS/'Cover_Letter_Draft.docx')

print(main_path);print(supp_path)
